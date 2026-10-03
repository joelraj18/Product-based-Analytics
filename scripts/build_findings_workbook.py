#!/usr/bin/env python3
"""Builds demo-data/WorkX_demo_data_and_findings.xlsx from FINDINGS.md,
demo-data/summary.json and the exported CSVs.

Run `npm run export-demo` first, then:
    python3 scripts/build_findings_workbook.py
Requires openpyxl (pip install openpyxl).
"""
import csv
import gzip
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'demo-data'
OUT = DATA / 'WorkX_demo_data_and_findings.xlsx'

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', fgColor='1E293B')
TITLE_FONT = Font(name=FONT, size=14, bold=True, color='0F172A')
SUB_FONT = Font(name=FONT, size=10, italic=True, color='52514E')
HEAD_FONT = Font(name=FONT, size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name=FONT, size=10)
BOLD_FONT = Font(name=FONT, size=10, bold=True)
TOTAL_FILL = PatternFill('solid', fgColor='F1F5F9')
THIN = Side(style='thin', color='E1E0D9')
SEVERITY_FILL = {
    'Critical': PatternFill('solid', fgColor='FDE2E2'),
    'High': PatternFill('solid', fgColor='FEF3C7'),
    'Medium': PatternFill('solid', fgColor='E0EDFB'),
    'Low': PatternFill('solid', fgColor='F1F5F9'),
}

INR = '₹#,##0;(₹#,##0);-'
INR2 = '₹#,##0.00;(₹#,##0.00);-'
INT = '#,##0;(#,##0);-'
DEC1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'


def read_csv(name):
    opener = gzip.open if name.endswith('.gz') else open
    with opener(DATA / name, 'rt', newline='', encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        for k, v in r.items():
            if v == '':
                r[k] = None
            elif re.fullmatch(r'-?\d+', v or '') and not (len(v) > 1 and v.startswith('0')):
                r[k] = int(v)
            elif re.fullmatch(r'-?\d+\.\d+', v or ''):
                r[k] = float(v)
    return rows


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def autosize(ws, min_w=8, max_w=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            if text.startswith('='):
                text = '0000000000'
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), max(len(t) for t in text.split('\n')))
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def number_format_for(col):
    c = col.lower()
    if 'pct' in c:
        return '0.0'
    if any(k in c for k in ('cost', 'op1', 'op2')) and 'volume' not in c:
        return INR2 if 'per_contact' in c or 'per_unit' in c else INR
    if any(k in c for k in ('fte', 'hc', 'shrinkage', 'sl_', 'occupancy', 'impact')):
        return DEC1
    return INT


def data_sheet(wb, title, rows, note=None, formats=None):
    ws = wb.create_sheet(title)
    cols = list(rows[0].keys()) if rows else []
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = SUB_FONT
        start = 3
    for j, c in enumerate(cols, 1):
        ws.cell(row=start, column=j, value=c)
    style_header(ws, start, len(cols))
    for i, r in enumerate(rows, start + 1):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=i, column=j, value=r[c])
            cell.font = BODY_FONT
            if isinstance(r[c], (int, float)):
                cell.number_format = (formats or {}).get(c, number_format_for(c))
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    if rows:
        ws.auto_filter.ref = f'A{start}:{get_column_letter(len(cols))}{start + len(rows)}'
    autosize(ws)
    return ws, start


def parse_findings():
    md = (ROOT / 'FINDINGS.md').read_text(encoding='utf-8')
    audit = []
    part1 = md.split('## Part 1')[1].split('## Part 2')[0]
    for line in part1.splitlines():
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if len(cells) == 6 and cells[0].isdigit():
            audit.append([int(cells[0])] + [re.sub(r'[`*]', '', c) for c in cells[1:]])
    actions = []
    sec = md.split('### 2.8 Recommended actions')[1].split('---')[0]
    for line in sec.splitlines():
        m = re.match(r'\d+\.\s+(.*)', line.strip())
        if m:
            actions.append(re.sub(r'[`*]', '', m.group(1)))
    snapshot = re.search(r'Snapshot exported on \*\*(.+?)\*\*', md)
    return audit, actions, snapshot.group(1) if snapshot else ''


def findings_sheet(wb, summary):
    audit, actions, snapshot = parse_findings()
    ws = wb.active
    ws.title = 'Findings'
    ws['A1'] = 'WorkX — Audit & Planning Findings'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = (f"Snapshot {snapshot} · plan horizon wk {summary['plan_horizon']['first_week']} → "
                f"{summary['plan_horizon']['last_week']} ({summary['plan_horizon']['weeks']} weeks) · "
                f"currency {summary['currency']} · source: FINDINGS.md, demo-data/summary.json")
    ws['A2'].font = SUB_FONT

    ws['A4'] = 'Part 1 — Audit findings (bugs found and fixed)'
    ws['A4'].font = BOLD_FONT
    head = ['#', 'Severity', 'Area', 'Symptom', 'Root cause', 'Fix']
    for j, h in enumerate(head, 1):
        ws.cell(row=5, column=j, value=h)
    style_header(ws, 5, len(head))
    r = 6
    for row in audit:
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=THIN)
        ws.cell(row=r, column=2).fill = SEVERITY_FILL.get(row[1], PatternFill())
        r += 1
    ws.cell(row=r, column=1, value='Count by severity').font = BOLD_FONT
    r += 1
    for sev in ['Critical', 'High', 'Medium', 'Low']:
        ws.cell(row=r, column=2, value=sev).font = BODY_FONT
        c = ws.cell(row=r, column=3, value=f'=COUNTIF($B$6:$B${5 + len(audit)},B{r})')
        c.font = BODY_FONT
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='Part 2 — Key planning insights').font = BOLD_FONT
    r += 1
    t = summary['totals']
    occ = '; '.join(f"{o['line_id']} {o['erlang_implied_occupancy_pct']}% vs plan {o['planned_occupancy_pct']}%" for o in summary['occupancy_check'])
    sc = summary['scenario_volume_plus20']
    k = summary['kpis_last_13_weeks']
    top = summary['defect_pareto'][:3]
    insights = [
        ('Demand & peak', f"Required FTE {t['required_fte_week1']} in week 1 → peak {t['peak_required_fte']}; peak weeks {', '.join(t['peak_weeks'])}."),
        ('Hiring', f"{t['perm_hires']} permanent + {t['temp_hires']} seasonal hires; the peak sits at the 7-week hiring lead time, so classes must start now."),
        ('Cost vs OP', f"Plan cost {t['variable_cost_fmt']}; {t['vs_op2_pct']:+}% vs OP2, {t['vs_op1_pct']:+}% vs OP1; cost per contact ₹{t['cost_per_contact']}."),
        ('Occupancy realism', f"Erlang-implied peak occupancy: {occ}. Plan understates real-time staffing."),
        ('Scenario +20% volume', f"With current hiring: {sc['current_hiring_plan']['weeks_short_after_ot']} weeks short after max OT; re-planned: {sc['replanned']['weeks_short_after_ot']} weeks (inside lead time)."),
        ('KPIs (13 wks)', f"WAPE {k['wape_pct']}%, SL {k['service_level_pct']}% vs {k['sl_target_pct']}% ({k['sl_weeks_met_pct']}% of line-weeks met), HC adherence {k['hc_adherence_pct']}%."),
        ('Root causes', '; '.join(f"{p['category']} {p['pct']}%" for p in top) + f" (cumulative {top[-1]['cumPct']}%)."),
    ]
    for label, text in insights:
        ws.cell(row=r, column=1, value=label).font = BOLD_FONT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=text)
        c.font = BODY_FONT
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='Recommended actions').font = BOLD_FONT
    r += 1
    for i, a in enumerate(actions, 1):
        ws.cell(row=r, column=1, value=i).font = BODY_FONT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=a)
        c.font = BODY_FONT
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 30
        r += 1

    for col, w in zip('ABCDEF', [22, 11, 18, 45, 45, 45]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A6'


def summary_sheet(wb, summary):
    ws = wb.create_sheet('Summary KPIs')
    ws['A1'] = 'Headline numbers (exported values)'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = 'Values exported by `npm run export-demo` from the app\'s planning engine; % values are percentages (not fractions).'
    ws['A2'].font = SUB_FONT
    rows = []
    t = summary['totals']
    rows += [('Plan', 'Required FTE week 1', t['required_fte_week1'], DEC1),
             ('Plan', 'Peak required FTE', t['peak_required_fte'], DEC1),
             ('Plan', 'Permanent hires', t['perm_hires'], INT),
             ('Plan', 'Seasonal temp hires', t['temp_hires'], INT),
             ('Plan', 'OT hours', t['ot_hours'], INT),
             ('Cost', 'Variable cost (horizon)', round(t['variable_cost']), INR),
             ('Cost', 'Variance vs OP1 (%)', t['vs_op1_pct'], '0.0'),
             ('Cost', 'Variance vs OP2 (%)', t['vs_op2_pct'], '0.0'),
             ('Cost', 'Cost per contact', t['cost_per_contact'], INR2)]
    for key, label in [('wape_pct', 'Forecast WAPE (%)'), ('bias_pct', 'Forecast bias (%)'), ('service_level_pct', 'Service level (%)'),
                       ('sl_target_pct', 'SL target (%)'), ('sl_weeks_met_pct', 'Line-weeks meeting SL (%)'), ('occupancy_pct', 'Occupancy (%)'),
                       ('hc_adherence_pct', 'HC plan adherence (%)'), ('shrinkage_vs_plan_pts', 'Shrinkage vs plan (pts)'), ('cost_vs_plan_pct', 'Cost vs plan (%)')]:
        rows.append(('KPIs (last 13 wks)', label, summary['kpis_last_13_weeks'][key], '0.0'))
    for o in summary['occupancy_check']:
        rows.append(('Erlang check', f"{o['line_id']} planned occupancy (%)", o['planned_occupancy_pct'], '0.0'))
        rows.append(('Erlang check', f"{o['line_id']} Erlang-implied peak occupancy (%)", o['erlang_implied_occupancy_pct'], '0.0'))
    for j, h in enumerate(['Group', 'Metric', 'Value'], 1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, 3)
    for i, (g, m, v, fmt) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=g).font = BODY_FONT
        ws.cell(row=i, column=2, value=m).font = BODY_FONT
        c = ws.cell(row=i, column=3, value=v)
        c.font = BODY_FONT
        c.number_format = fmt
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 18
    ws.freeze_panes = 'A5'


def budget_sheet(wb):
    rows = read_csv('budget_monthly.csv')
    for r in rows:  # variances become live formulas below
        r.pop('var_vs_op1_pct', None)
        r.pop('var_vs_op2_pct', None)
        r.pop('cost_per_contact', None)
    ws, start = data_sheet(wb, 'Budget vs OP', rows, note='Monthly variable cost of the plan vs OP1/OP2. Cost per contact and variances are formulas; edit OP cells (blue) to test targets.')
    cols = list(rows[0].keys())
    idx = {c: get_column_letter(cols.index(c) + 1) for c in cols}
    extra = ['cost_per_contact', 'var_vs_op1', 'var_vs_op2']
    base_col = len(cols)
    for j, h in enumerate(extra, base_col + 1):
        ws.cell(row=start, column=j, value=h)
    style_header(ws, start, base_col + len(extra))
    first, last = start + 1, start + len(rows)
    V, T, O1, O2 = idx['volume'], idx['total_cost'], idx['op1_cost'], idx['op2_cost']
    for r in range(first, last + 1):
        for c in (O1, O2):
            ws[f'{c}{r}'].font = Font(name=FONT, size=10, color='0000FF')
        ws.cell(row=r, column=base_col + 1, value=f'=IF({V}{r}>0,{T}{r}/{V}{r},0)').number_format = INR2
        ws.cell(row=r, column=base_col + 2, value=f'=IF({O1}{r}>0,{T}{r}/{O1}{r}-1,0)').number_format = PCT
        ws.cell(row=r, column=base_col + 3, value=f'=IF({O2}{r}>0,{T}{r}/{O2}{r}-1,0)').number_format = PCT
        for j in range(base_col + 1, base_col + 4):
            ws.cell(row=r, column=j).font = BODY_FONT
    tr = last + 1
    ws.cell(row=tr, column=1, value='Total').font = BOLD_FONT
    for c in cols[1:]:
        L = idx[c]
        cell = ws[f'{L}{tr}']
        cell.value = f'=SUM({L}{first}:{L}{last})'
        cell.font = BOLD_FONT
        cell.number_format = number_format_for(c)
    ws.cell(row=tr, column=base_col + 1, value=f'=IF({V}{tr}>0,{T}{tr}/{V}{tr},0)').number_format = INR2
    ws.cell(row=tr, column=base_col + 2, value=f'=IF({O1}{tr}>0,{T}{tr}/{O1}{tr}-1,0)').number_format = PCT
    ws.cell(row=tr, column=base_col + 3, value=f'=IF({O2}{tr}>0,{T}{tr}/{O2}{tr}-1,0)').number_format = PCT
    for j in range(1, base_col + 4):
        ws.cell(row=tr, column=j).fill = TOTAL_FILL
        ws.cell(row=tr, column=j).font = BOLD_FONT
    autosize(ws)


def scenario_sheet(wb):
    rows = read_csv('scenario_volume_plus20.csv')
    ws, start = data_sheet(wb, 'Scenario +20% volume', rows, note='Baseline vs +20% volume, holding the current hiring plan, and after re-planning hires. Deltas are formulas.',
                           formats={'baseline': '#,##0.0', 'scenario_current_hiring_plan': '#,##0.0', 'scenario_replanned': '#,##0.0'})
    n = len(rows[0])
    for j, h in enumerate(['delta_current_plan', 'delta_replanned'], n + 1):
        ws.cell(row=start, column=j, value=h)
    style_header(ws, start, n + 2)
    for r in range(start + 1, start + len(rows) + 1):
        ws.cell(row=r, column=n + 1, value=f'=C{r}-B{r}').number_format = '#,##0.0;(#,##0.0);-'
        ws.cell(row=r, column=n + 2, value=f'=D{r}-B{r}').number_format = '#,##0.0;(#,##0.0);-'
        ws.cell(row=r, column=n + 1).font = BODY_FONT
        ws.cell(row=r, column=n + 2).font = BODY_FONT
    autosize(ws)


def main():
    summary = json.loads((DATA / 'summary.json').read_text(encoding='utf-8'))
    wb = Workbook()
    findings_sheet(wb, summary)
    summary_sheet(wb, summary)
    data_sheet(wb, 'Capacity Plan', read_csv('capacity_plan_weekly.csv'), note='Weekly headcount plan, all lines × horizon (recommended hiring).')
    budget_sheet(wb)
    data_sheet(wb, 'Forecast Accuracy', read_csv('forecast_accuracy.csv'), note='26-week backtest of the forecast method in use per line.')
    scenario_sheet(wb)
    data_sheet(wb, 'Actuals', read_csv('actuals_weekly.csv'), note='Weekly actuals by line (last 26 weeks).')
    data_sheet(wb, 'Defects', read_csv('defects.csv'))
    data_sheet(wb, 'Risks', read_csv('risks.csv'))
    data_sheet(wb, 'Plan Lines', read_csv('plan_lines.csv'), note='Planning assumptions per program/queue (AHT in seconds, rates in %).')
    data_sheet(wb, 'Events', read_csv('events.csv'))
    data_sheet(wb, 'OP Targets', read_csv('op_targets.csv'))
    data_sheet(wb, 'Orders', read_csv('orders.csv.gz'))
    data_sheet(wb, 'Volume History', read_csv('volume_history_daily.csv'), note='Daily contacts and AHT per line (3 years). Re-importable in Demand Forecast.')
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title not in ('Findings', 'Summary KPIs')
    # openpyxl stores formulas without cached results; make Excel/LibreOffice
    # compute everything when the file is opened.
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f'Wrote {OUT.relative_to(ROOT)} with sheets: {", ".join(wb.sheetnames)}')


if __name__ == '__main__':
    main()
